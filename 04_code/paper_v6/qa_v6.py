"""Submission checks for editorial v6 against frozen v5."""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
OLD = ROOT / "11_delivery/v5/F_final_candidate_v5_参考论文深度吸收完整版.docx"
STEM = "F_final_candidate_v6_获奖范文格式语言优化版"
DOCX = ROOT / "11_delivery/v6" / (STEM + ".docx")
PDF = ROOT / "11_delivery/v6" / (STEM + ".pdf")
MANIFEST = ROOT / "08_paper/v6/BUILD_MANIFEST_v6.json"
REPORT = ROOT / "10_review/paper_v6/FINAL_QA_v6.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def media_hashes(path: Path) -> Counter:
    with ZipFile(path) as docx:
        return Counter(hashlib.sha256(docx.read(name)).hexdigest()
                       for name in docx.namelist() if name.startswith("word/media/"))


def math_text(doc) -> list[str]:
    texts = [''.join(m.itertext()) for p in doc.paragraphs
             for m in p._p.xpath('.//m:oMath')]
    return [s.replace('−', '-').replace('ℓ', 'l') for s in texts]


def table_is_three_line(table) -> bool:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        return False
    present = {child.tag.rsplit('}', 1)[-1] for child in borders}
    if not {'top', 'bottom'} <= present or present & {'left', 'right', 'insideH', 'insideV'}:
        return False
    header = table.rows[0]._tr.get_or_add_trPr()
    if header.find(qn("w:tblHeader")) is None:
        return False
    for row in table.rows:
        if row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is None:
            return False
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            if tcpr.find(qn("w:shd")) is not None:
                return False
            cb = tcpr.find(qn("w:tcBorders"))
            if cb is not None:
                names = {child.tag.rsplit('}', 1)[-1] for child in cb}
                if names & {'left', 'right', 'insideH', 'insideV'}:
                    return False
    return True


def main() -> None:
    old = Document(OLD)
    new = Document(DOCX)
    pdf = PdfReader(PDF)
    texts = [(page.extract_text() or '') for page in pdf.pages]
    flat = re.sub(r"\s+", "", ''.join(texts))

    old_tables = [[cell.text for row in table.rows for cell in row.cells]
                  for table in old.tables]
    new_tables = [[cell.text for row in table.rows for cell in row.cells]
                  for table in new.tables]
    captions = [p.text for p in new.paragraphs if p.style
                and p.style.name == "PaperCaption" and re.match(r"^图\d+\s", p.text)]
    fig_numbers = [int(re.match(r"^图(\d+)", text).group(1)) for text in captions]
    heading_pages = [set(line.strip() for line in text.splitlines()) for text in texts]
    toc = []
    for p in new.paragraphs:
        if p.style and p.style.name == "PaperTOC":
            title, printed = p.text.split("\t")
            found = [i + 1 for i in range(3, len(texts)) if title in heading_pages[i]]
            toc.append({"title": title, "printed": int(printed), "found": found})
    with ZipFile(DOCX) as z:
        thumbnail_parts = [x for x in z.namelist() if "thumbnail" in x.lower()]
        thumbnail_rel = b"metadata/thumbnail" in z.read("_rels/.rels")
    body = '\n'.join(p.text for p in new.paragraphs if p.style and p.style.name == "PaperBody")
    cited = {int(n) for n in re.findall(r"\[([1-7])\]", body)}
    refs = {int(m.group(1)) for p in new.paragraphs if p.style and p.style.name == "PaperRef"
            if (m := re.match(r"\[(\d+)\]", p.text))}
    key_values = ["0.2846", "0.2278", "3.108", "2.3452", "2.2672",
                  "0.33997658", "0.27987813", "0.000146888", "0.000111642",
                  "0.000114134", "0.361995", "9.325359", "2.300064",
                  "17.0233", "82.9767", "79.76", "87.13", "77.93",
                  "83.48", "76.11", "79.84"]
    missing_values = [value for value in key_values if value not in flat]

    checks = {
        "source_v5_unchanged": sha(OLD) == "698de15a867876e22b9194ef03f0cb6110b9293d1bdd841b077c491b0b683ecc",
        "pdf_opens_43_pages": len(pdf.pages) == 43 and not pdf.is_encrypted,
        "pdf_no_blank_pages": all(len(text) >= 70 for text in texts),
        "tables_content_preserved": old_tables == new_tables,
        "all_23_academic_tables_three_line": len(new.tables) == 24 and all(table_is_three_line(t) for t in new.tables[1:]),
        "original_29_media_preserved": not (media_hashes(OLD) - media_hashes(DOCX)) and len(new.inline_shapes) == 33,
        "math_content_preserved": math_text(old) == math_text(new),
        "figures_1_to_29_sequential": fig_numbers == list(range(1, 30)),
        "four_route_diagrams_present": all(any(f"问题{i}技术路线" in c for c in captions)
                                           for i in ("一", "二", "三", "四")),
        "toc_37_entries_match_final_pdf": len(toc) == 37 and all(row["found"] == [row["printed"] + 1] for row in toc),
        "citation_list_closed": cited == refs == set(range(1, 8)),
        "key_result_values_present": not missing_values,
        "no_thumbnail_or_dangling_relationship": not thumbnail_parts and not thumbnail_rel,
        "anonymous_core_properties": not new.core_properties.author and not new.core_properties.last_modified_by,
        "no_reference_advertisement": "anjia211014" not in flat and "代充服务" not in flat,
        "single_ai_usage_notice": body.count("AI辅助使用说明") == 1,
        "fixed_equation_reference": "式（9）—（9）" not in body,
    }
    report = {
        "pass": all(checks.values()), "checks": checks,
        "pdf_physical_pages": len(pdf.pages), "pdf_printed_last_page": len(pdf.pages) - 1,
        "toc": toc, "figures": len(captions), "tables_academic": len(new.tables) - 1,
        "math_objects": len(math_text(new)), "embedded_shapes": len(new.inline_shapes),
        "missing_key_values": missing_values,
        "docx_sha256": sha(DOCX), "pdf_sha256": sha(PDF),
        "officecli_schema_validation": "previous binary: success, 0 warnings; current binary missing System.IO.Pipelines",
        "visual_review": "PDF rounds 1, 2, 3 and final round 4; every final physical page reviewed in contact sheets; key pages enlarged",
        "layout_note": "Physical pages 25 and 27 retain figure-led whitespace to preserve graph label readability.",
    }
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    manifest.update({"docx_sha256": report["docx_sha256"], "pdf_sha256": report["pdf_sha256"],
                     "pdf_physical_pages": report["pdf_physical_pages"], "qa_pass": report["pass"]})
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"pass": report["pass"], "failed": [k for k, v in checks.items() if not v],
                      "pages": len(pdf.pages), "figures": len(captions),
                      "academic_tables": len(new.tables) - 1}, ensure_ascii=False))
    if not report["pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
