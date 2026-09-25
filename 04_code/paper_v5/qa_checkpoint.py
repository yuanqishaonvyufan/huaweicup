"""Structural QA for v5 DOCX/PDF checkpoints; no model computation."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "08_paper/v5"
OUT = ROOT / "11_delivery/v5"
STEM = "F_final_candidate_v5_参考论文深度吸收完整版"


def numbered(doc: Document, prefix: str) -> list[int]:
    return [int(m.group(1)) for p in doc.paragraphs if (m := re.match(rf"^{prefix}(\d+) ", p.text))]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", choices=["q1", "q2", "q3", "q4", "final"], required=True)
    opt = parser.parse_args()
    name = STEM if opt.stage == "final" else f"F_v5_{opt.stage.upper()}_checkpoint"
    docx, pdf = OUT / (name + ".docx"), OUT / (name + ".pdf")
    manifest = json.loads((SRC / "BUILD_MANIFEST_v5.json").read_text(encoding="utf-8"))
    assert manifest["stage"] == opt.stage
    assert docx.stat().st_size > 100_000 and pdf.stat().st_size > 100_000
    doc = Document(docx)
    figures, tables = numbered(doc, "图"), numbered(doc, "表")
    equations = [int(m.group(1)) for p in doc.paragraphs
                 if p.style is not None and p.style.name == "PaperEquation"
                 if (m := re.search(r"（(\d+)）", p.text))]
    assert figures == list(range(1, manifest["figures"] + 1)), figures
    assert tables == list(range(1, manifest["tables_body"] + 1)), tables
    assert equations == list(range(1, manifest["equations"] + 1)), equations
    assert all(manifest["baseline_media_preserved"].values())
    assert len(doc.inline_shapes) == 13 + (len(figures) - 9)
    assert not any(p.text.strip().startswith(("@@BEFORE", "|")) for p in doc.paragraphs)
    assert not any("图#" in p.text or "表#" in p.text or "（#）" in p.text for p in doc.paragraphs)
    pages = len(PdfReader(pdf).pages)
    assert pages >= 34
    print(json.dumps({"stage": opt.stage, "pages": pages, "figures": len(figures),
                      "body_tables": len(tables), "numbered_equations": len(equations),
                      "baseline_media_preserved": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
