"""Measure reference audit, user-designated v3, and final v5 consistently."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
STEM = "F_final_candidate_v5_参考论文深度吸收完整版"
V3_PDF = ROOT / "tmp/paper_v5/render_v3/baseline_v3_user.pdf"
V5_PDF = ROOT / "11_delivery/v5" / (STEM + ".pdf")
V3_DOCX = ROOT / "08_paper/v5/baseline_v3_user.docx"
V5_DOCX = ROOT / "11_delivery/v5" / (STEM + ".docx")
OUT = ROOT / "10_review/paper_v5/REFERENCE_VS_V5_METRICS.md"
HAN = re.compile(r"[\u3400-\u9fff]")
KEYS = ["q1", "q2", "q3", "q4", "validation", "sensitivity", "evaluation", "improvement", "references"]
HEADINGS = {
    "q1": "5.1 问题一的模型的建立和求解",
    "q2": "5.2 问题二的模型的建立和求解",
    "q3": "5.3 问题三的模型的建立和求解",
    "q4": "5.4 问题四的模型的建立和求解",
    "validation": "六、模型检验与分析",
    "sensitivity": "6.2 灵敏度分析",
    "evaluation": "七、模型的评价",
    "improvement": "八、模型的改进与推广",
    "references": "参考文献",
}


def count_fig_table_eq(docx: Path) -> tuple[int, int, int, int]:
    doc = Document(docx)
    figures = [int(m.group(1)) for p in doc.paragraphs if (m := re.match(r"^图(\d+) ", p.text))]
    tables = [int(m.group(1)) for p in doc.paragraphs if (m := re.match(r"^表(\d+) ", p.text))]
    appendix = [p.text for p in doc.paragraphs if re.match(r"^表[A-Z]\d+ ", p.text)]
    equations = [int(m.group(1)) for p in doc.paragraphs
                 if p.style is not None and p.style.name == "PaperEquation"
                 if (m := re.search(r"（(\d+)）", p.text))]
    assert sorted(figures) == list(range(1, len(figures) + 1))
    assert sorted(tables) == list(range(1, len(tables) + 1))
    assert sorted(equations) == list(range(1, len(equations) + 1))
    return len(figures), len(tables), len(appendix), len(equations)


def measure(pdf: Path, docx: Path) -> dict:
    reader = PdfReader(pdf)
    texts = [(p.extract_text() or "") for p in reader.pages]
    body = "\n".join(texts[3:])
    positions = {key: body.find(heading) for key, heading in HEADINGS.items()}
    if any(value < 0 for value in positions.values()):
        raise ValueError(f"Missing heading in {pdf}: {positions}")
    assert all(positions[KEYS[i]] < positions[KEYS[i + 1]] for i in range(len(KEYS) - 1))
    intro = body.find("一、引言")
    assert 0 <= intro < positions["q1"]
    section_end = {"q1": "q2", "q2": "q3", "q3": "q4", "q4": "validation",
                   "validation": "evaluation", "sensitivity": "evaluation",
                   "evaluation": "improvement", "improvement": "references"}
    sections = {key: len(HAN.findall(body[positions[key]:positions[section_end[key]]])) for key in section_end}
    sections["validation_without_sensitivity"] = len(HAN.findall(body[positions["validation"]:positions["sensitivity"]]))
    refs_page = next(i + 1 for i, text in enumerate(texts[3:], 3) if "参考文献" in text.splitlines())
    body_touched = refs_page - 3
    f, t, appendix, eq = count_fig_table_eq(docx)
    return {
        "pages": len(texts), "body_touched_pages": body_touched,
        "all_han": sum(len(HAN.findall(text)) for text in texts),
        "body_han_excluding_refs": len(HAN.findall(body[intro:positions["references"]])),
        "figures": f, "body_tables": t, "appendix_tables": appendix,
        "total_tables": t + appendix, "equations": eq,
        "sections": sections, "references_physical_page": refs_page,
    }


def main() -> None:
    v3, v5 = measure(V3_PDF, V3_DOCX), measure(V5_PDF, V5_DOCX)
    ref = {"pages": 29, "body_touched_pages": 22, "all_han": 11498,
           "body_han_excluding_refs": 9736, "figures": 20, "body_tables": 10,
           "appendix_tables": 1, "total_tables": 11, "equations": 54,
           "sections": {"q1": 2057, "q2": 1613, "q3": 1483, "q4": 1535,
                        "validation": 461, "sensitivity": None,
                        "evaluation": 374, "improvement": 441}}
    def row(label, key, nested=False):
        def get(x):
            value = x["sections"].get(key, "未单列") if nested else x[key]
            return "未单列" if value is None else value
        return f"|{label}|{get(ref)}|{get(v3)}|{get(v5)}|"
    table = ["|项目|Reference|v3|v5|", "|---|---:|---:|---:|"]
    for label, key in [
        ("总页数", "pages"), ("正文触及页数", "body_touched_pages"),
        ("全稿中文字符", "all_han"), ("正文中文字符（不含参考文献）", "body_han_excluding_refs"),
        ("编号公式", "equations"), ("正文 Figure", "figures"),
        ("Table 总数", "total_tables"),
        ("正文 Table", "body_tables"), ("附录 Table", "appendix_tables"),
    ]:
        table.append(row(label, key))
    for label, key in [
        ("Q1 中文字符", "q1"), ("Q2 中文字符", "q2"),
        ("Q3 中文字符", "q3"), ("Q4 中文字符", "q4"),
        ("模型检验与灵敏度 中文字符", "validation"),
        ("灵敏度小节 中文字符", "sensitivity"),
        ("模型评价 中文字符", "evaluation"),
        ("改进与推广 中文字符", "improvement"),
    ]:
        table.append(row(label, key, nested=True))
    report = "# Reference versus v3 versus v5 metrics\n\n" + "\n".join(table) + "\n\n"
    report += ("口径：Reference 来自原参考 PDF 逐页结构审计（29页、20图、11表、54式）；"
               "v3 是用户指定 DOCX 原字节副本的本环境只读重渲染；v5 是最终交付 PDF。"
               "中文字符仅数文本层 U+3400–U+9FFF，不含图内不可抽取文字。"
               "正文触及页数从引言首个物理页计至参考文献起始物理页，若该页共享末段正文则计入；"
               "Reference 正文22页按原逐页审计口径。正文 Figure/Table 不计附录表。\n\n")
    report += (f"v3 参考文献起于物理第 {v3['references_physical_page']} 页；"
               f"v5 起于物理第 {v5['references_physical_page']} 页。"
               "v3 原审计报告以另一 PDF 渲染测得约 18,710 个全稿中文字符；"
               "上表对 v3 和 v5 均采用本次同一提取器，因此同列差异可直接比较。\n")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(report, encoding="utf-8")
    print({"v3": v3, "v5": v5})


if __name__ == "__main__":
    main()
