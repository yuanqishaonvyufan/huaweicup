"""Final v5 structural, source-hash, TOC, and privacy checks."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
STEM = "F_final_candidate_v5_参考论文深度吸收完整版"
DOCX = ROOT / "11_delivery/v5" / (STEM + ".docx")
PDF = ROOT / "11_delivery/v5" / (STEM + ".pdf")
SRC = ROOT / "08_paper/v5"
OUT = ROOT / "10_review/paper_v5/FINAL_QA_v5.json"
RUNS = {
    "SUPP-Q1-VIS-20260925-v1": "supplementary_q1_vis.py",
    "SUPP-Q2-VIS-20260925-v1": "supplementary_q2_vis.py",
    "SUPP-Q3-VIS-20260925-v1": "supplementary_q3_vis.py",
    "SUPP-Q4-VIS-20260925-v1": "supplementary_q4_vis.py",
}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    manifest_path = SRC / "BUILD_MANIFEST_v5.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["stage"] == "final" and manifest["toc_refreshed"] and manifest["toc_rows"] == 37
    assert sha(DOCX) == manifest["docx_sha256"]
    assert all(manifest["baseline_media_preserved"].values())
    doc = Document(DOCX)
    reader = PdfReader(PDF)
    pages = [(page.extract_text() or "") for page in reader.pages]
    assert len(pages) == 52
    assert min(len(text.strip()) for text in pages[3:]) > 80
    lines = [set(line.strip() for line in text.splitlines()) for text in pages]
    toc = []
    for p in doc.paragraphs:
        if p.style is None or p.style.name != "PaperTOC":
            continue
        title, printed = p.text.split("\t")
        found = [i + 1 for i, page_lines in enumerate(lines[3:], 3) if title in page_lines]
        assert found == [int(printed) + 1], (title, printed, found)
        toc.append(title)
    assert len(toc) == 37
    captions = [p.text for p in doc.paragraphs]
    figures = [int(m.group(1)) for text in captions if (m := re.match(r"^图(\d+) ", text))]
    tables = [int(m.group(1)) for text in captions if (m := re.match(r"^表(\d+) ", text))]
    appendix_tables = [text.split()[0] for text in captions if re.match(r"^表[A-Z]\d+ ", text)]
    equations = [int(m.group(1)) for p in doc.paragraphs
                 if p.style is not None and p.style.name == "PaperEquation"
                 if (m := re.search(r"（(\d+)）", p.text))]
    assert figures == list(range(1, 26))
    assert tables == list(range(1, 21))
    assert appendix_tables == ["表A1", "表B1", "表D1"]
    assert equations == list(range(1, 29))
    assert len(doc.inline_shapes) == 29
    appendix_text = "\n".join(p.text for p in doc.paragraphs if p.text.startswith("附录"))
    whole_text = "\n".join(p.text for p in doc.paragraphs)
    assert not re.search(r"(?:C:|D:)\\|github\.com/yuanqishaonvyufan|(?:^|\s)(?:Round|Gate|checkpoint|commit hash)(?:\s|$)", whole_text, flags=re.I | re.M)
    assert not re.search(r"74\.3|87\.9|12%/88%", whole_text)
    source_runs = {}
    for run_id, script in RUNS.items():
        path = ROOT / "06_results/raw" / run_id / "manifest.json"
        run = json.loads(path.read_text(encoding="utf-8"))
        assert run["run_id"] == run_id
        assert sha(ROOT / "04_code/paper_v5" / script) == run["code_sha256"]
        for category in ("inputs_sha256", "outputs_sha256"):
            assert all(sha(ROOT / name) == expected for name, expected in run[category].items())
        source_runs[run_id] = {"input_count": len(run["inputs_sha256"]),
                               "output_count": len(run["outputs_sha256"]), "hashes_match": True}
    manifest["pdf_sha256"] = sha(PDF)
    manifest["pdf_pages"] = len(pages)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    result = {"pass": True, "physical_pages": len(pages), "toc_entries_verified": len(toc),
              "figures": len(figures), "body_tables": len(tables), "appendix_tables": len(appendix_tables),
              "equations": len(equations), "embedded_shapes": len(doc.inline_shapes),
              "baseline_media_preserved": True, "supplementary_runs": source_runs,
              "docx_sha256": sha(DOCX), "pdf_sha256": sha(PDF)}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
