"""Expand the user-designated v3 DOCX with verified local evidence.

The source DOCX is preserved byte-for-byte in 08_paper/v5. This builder starts
from it on every run, so section checkpoints cannot compound edits.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert
from lxml import etree
from PIL import Image
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "08_paper/v5"
BASE = SRC / "baseline_v3_user.docx"
OUT = ROOT / "11_delivery/v5"
STEM = "F_final_candidate_v5_参考论文深度吸收完整版"
PHASES = [
    ("q1", "05_Q1_additions.md"),
    ("q2", "06_Q2_additions.md"),
    ("q3", "07_Q3_additions.md"),
    ("q4", "08_Q4_additions.md"),
    ("final", "09_validation_back_additions.md"),
]
XSL = Path("C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL")


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def find_anchor(doc: Document, key: str) -> Paragraph | None:
    if key == "END":
        return None
    hits = [p for p in doc.paragraphs if p.text.strip() == key and (p.style is None or p.style.name != "PaperTOC")]
    if len(hits) != 1:
        raise ValueError(f"Anchor {key!r}: expected one non-TOC paragraph, found {len(hits)}")
    return hits[0]


def add_para(doc: Document, anchor: Paragraph | None, text: str, style: str = "PaperBody") -> Paragraph:
    if anchor is None:
        p = doc.add_paragraph(style=style)
    else:
        element = OxmlElement("w:p")
        anchor._p.addprevious(element)
        p = Paragraph(element, anchor._parent)
        p.style = style
    if text:
        p.add_run(text)
    p.paragraph_format.widow_control = True
    return p


def add_figure(doc: Document, anchor: Paragraph | None, caption: str, source: str, sources: list[dict]) -> None:
    path = ROOT / source
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as im:
        ratio = im.height / im.width
    width = 15.5 if ratio < .8 else 13.7
    p = add_para(doc, anchor, "", "PaperCaption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    c = add_para(doc, anchor, "图# " + caption, "PaperCaption")
    c.paragraph_format.keep_together = True
    sources.append({"source": source, "sha256": sha(path), "caption": caption, "width_cm": width})


def add_equation(doc: Document, anchor: Paragraph | None, tex: str, xsl) -> None:
    p = add_para(doc, anchor, "", "PaperEquation")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    mathml = etree.fromstring(convert(tex).encode("utf-8"))
    p._p.append(xsl(mathml).getroot())
    p.add_run("    （#）")


def add_table(doc: Document, anchor: Paragraph | None, rows: list[list[str]]) -> None:
    n = len(rows[0])
    if not rows or any(len(row) != n for row in rows):
        raise ValueError("Malformed table")
    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if anchor is not None:
        anchor._p.addprevious(table._tbl)
    page = doc.sections[-1]
    total = page.page_width.cm - page.left_margin.cm - page.right_margin.cm
    weights = {3: [.24, .32, .44], 4: [.22, .24, .25, .29], 5: [.19, .19, .20, .20, .22]}.get(n, [1 / n] * n)
    for col, weight in zip(table.columns, weights):
        col.width = Cm(total * weight)
    for ri, row in enumerate(rows):
        cells = table.rows[0].cells if ri == 0 else table.add_row().cells
        trpr = table.rows[ri]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(total * weights[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = doc.styles["PaperTable"]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or len(value) > 18 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            run.font.size = Pt(10 if n >= 5 else 10.5)
            if ri == 0:
                run.bold = True
            tc = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EDF4" if ri == 0 else ("F8FAFC" if ri % 2 == 0 else "FFFFFF"))
            tc.append(shade)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + side)
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), "4")
                el.set(qn("w:color"), "D9D9D9")
                borders.append(el)
            tc.append(borders)
            margins = OxmlElement("w:tcMar")
            for side, value in (("top", "80"), ("bottom", "80"), ("left", "85"), ("right", "85")):
                el = OxmlElement("w:" + side)
                el.set(qn("w:w"), value)
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tc.append(margins)


def parse_and_insert(doc: Document, source: Path, figures: list[dict], xsl) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    anchor = None
    after_table = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("@@BEFORE "):
            anchor = find_anchor(doc, line[len("@@BEFORE "):])
            after_table = False
            continue
        if anchor is None and "@@BEFORE END" not in source.read_text(encoding="utf-8"):
            raise ValueError(f"No anchor before {source}:{i}")
        if line.startswith("!["):
            match = re.fullmatch(r"!\[(.+)\]\((.+)\)", line)
            if not match:
                raise ValueError(f"Invalid figure in {source}:{i}")
            add_figure(doc, anchor, match.group(1), match.group(2), figures)
        elif line.startswith("$$ ") and line.endswith(" $$"):
            add_equation(doc, anchor, line[3:-3].strip(), xsl)
        elif line.startswith("表# "):
            add_para(doc, anchor, line, "PaperCaption").paragraph_format.keep_with_next = True
        elif line.startswith("|"):
            rows = []
            while True:
                row = [v.strip() for v in line.strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", v) for v in row):
                    rows.append(row)
                if i >= len(lines) or not lines[i].strip().startswith("|"):
                    break
                line = lines[i].strip()
                i += 1
            add_table(doc, anchor, rows)
            after_table = True
        elif line.startswith("### "):
            p = add_para(doc, anchor, line[4:], "PaperH2")
            p.paragraph_format.space_before = Pt(6)
        else:
            p = add_para(doc, anchor, line)
            if after_table:
                p.paragraph_format.space_before = Pt(7)
                after_table = False


def replace_runs(paragraph: Paragraph, transform) -> None:
    for run in paragraph.runs:
        before = run.text
        after = transform(before)
        if before != after:
            run.text = after


def renumber(doc: Document, old_elements: set, count_name: str, prefix: str) -> tuple[dict[int, int], int]:
    mapping = {}
    count = 0
    pattern = re.compile(rf"^{prefix}(\d+|#) ")
    for p in doc.paragraphs:
        if p.style is None or p.style.name != "PaperCaption":
            continue
        match = pattern.match(p.text)
        if not match:
            continue
        count += 1
        if match.group(1) != "#":
            old = int(match.group(1))
            if old in mapping:
                raise ValueError(f"Duplicate old {count_name} {old}")
            mapping[old] = count
        p.text = pattern.sub(f"{prefix}{count} ", p.text, count=1)
    return mapping, count


def renumber_equations(doc: Document) -> tuple[dict[int, int], int]:
    mapping = {}
    count = 0
    pattern = re.compile(r"（(\d+|#)）")
    for p in doc.paragraphs:
        if p.style is None or p.style.name != "PaperEquation":
            continue
        hit = pattern.search(p.text)
        if not hit:
            continue
        count += 1
        if hit.group(1) != "#":
            old = int(hit.group(1))
            if old in mapping:
                raise ValueError(f"Duplicate old equation {old}")
            mapping[old] = count
        replace_runs(p, lambda value: pattern.sub(f"（{count}）", value))
    return mapping, count


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--through", choices=[phase for phase, _ in PHASES], required=True)
    opt = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document(BASE)
    old_elements = {p._p for p in doc.paragraphs}
    old_text = {p._p: p.text for p in doc.paragraphs}
    xsl = etree.XSLT(etree.parse(str(XSL)))
    figures = []
    sources = []
    for phase, filename in PHASES:
        source = SRC / filename
        if not source.is_file():
            raise FileNotFoundError(source)
        parse_and_insert(doc, source, figures, xsl)
        sources.append({"phase": phase, "source": filename, "sha256": sha(source)})
        if phase == opt.through:
            break
    fig_map, fig_count = renumber(doc, old_elements, "figure", "图")
    table_map, table_count = renumber(doc, old_elements, "table", "表")
    equation_map, equation_count = renumber_equations(doc)
    for p in doc.paragraphs:
        if p._p not in old_elements or (p.style is not None and p.style.name in ("PaperCaption", "PaperEquation", "PaperTOC")):
            continue
        def change(value: str) -> str:
            value = re.sub(r"图(\d+)", lambda m: "图" + str(fig_map.get(int(m.group(1)), int(m.group(1)))), value)
            value = re.sub(r"表(\d+)", lambda m: "表" + str(table_map.get(int(m.group(1)), int(m.group(1)))), value)
            value = re.sub(r"式（(\d+)）", lambda m: "式（" + str(equation_map.get(int(m.group(1)), int(m.group(1)))) + "）", value)
            return value
        replace_runs(p, change)
        if p.text != change(old_text[p._p]):
            raise ValueError(f"Cannot update split reference: {old_text[p._p][:100]}")
    doc.core_properties.title = "算力约束下大语言模型的质量评价与资源配置及能力前沿预测"
    for field in ("author", "last_modified_by", "comments", "keywords", "identifier"):
        setattr(doc.core_properties, field, "")
    out_stem = STEM if opt.through == "final" else f"F_v5_{opt.through.upper()}_checkpoint"
    docx = OUT / (out_stem + ".docx")
    doc.save(docx)
    with ZipFile(BASE) as old, ZipFile(docx) as new:
        original_media = [p for p in old.namelist() if p.startswith("word/media/")]
        media_preserved = {p: old.read(p) == new.read(p) for p in original_media}
    if not all(media_preserved.values()):
        raise ValueError("Baseline embedded media changed")
    manifest = {
        "stage": opt.through,
        "baseline": "08_paper/v5/baseline_v3_user.docx",
        "baseline_sha256": sha(BASE),
        "source_files": sources,
        "figures": fig_count,
        "tables_body": table_count,
        "equations": equation_count,
        "figure_sources_new": figures,
        "old_figure_number_map": fig_map,
        "old_table_number_map": table_map,
        "old_equation_number_map": equation_map,
        "baseline_media_preserved": media_preserved,
        "docx_sha256": sha(docx),
    }
    (SRC / "BUILD_MANIFEST_v5.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"docx": str(docx), "stage": opt.through, "figures": fig_count,
                      "tables_body": table_count, "equations": equation_count}, ensure_ascii=False))


if __name__ == "__main__":
    main()
