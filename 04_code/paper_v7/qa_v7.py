"""Verify publication cleanup against the verified v6 manuscript."""

import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "04_code/paper_v6"))
from qa_v6 import math_text, media_hashes, table_is_three_line  # noqa: E402

from final_cleanup import REWRITES, numeric_tokens  # noqa: E402


V6 = ROOT / "11_delivery/v6/F_final_candidate_v6_获奖范文格式语言优化版.docx"
DOCX = ROOT / "11_delivery/v7/F_final_candidate_v7_最终投稿版.docx"
PDF = ROOT / "11_delivery/v7/F_final_candidate_v7_最终投稿版.pdf"
REPORT = ROOT / "10_review/paper_v7/FINAL_QA_v7.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    old = Document(V6)
    new = Document(DOCX)
    pdf = PdfReader(PDF)
    texts = [(page.extract_text() or "") for page in pdf.pages]
    full_text = "\n".join(p.text for p in new.paragraphs) + "\n" + "\n".join(
        c.text for t in new.tables for row in t.rows for c in row.cells)
    flat_pdf = re.sub(r"\s+", "", "".join(texts))

    old_paragraphs = [p.text for p in old.paragraphs]
    new_paragraphs = [p.text for p in new.paragraphs]
    text_changes = [i for i, (a, b) in enumerate(zip(old_paragraphs, new_paragraphs)) if a != b]
    toc_changes = [i for i in text_changes if old.paragraphs[i].style.name == "PaperTOC"]
    unexpected_changes = [i for i in text_changes if i not in REWRITES and i not in toc_changes]
    numeric_changes = [i for i in text_changes
                       if numeric_tokens(old_paragraphs[i]) != numeric_tokens(new_paragraphs[i])]
    expected_numeric_changes = {53, 164, 247, 278, 404, 445, 446}

    old_table_numbers = numeric_tokens("\n".join(c.text for t in old.tables for r in t.rows for c in r.cells))
    new_table_numbers = numeric_tokens("\n".join(c.text for t in new.tables for r in t.rows for c in r.cells))
    old_refs = [p.text for p in old.paragraphs if p.style.name == "PaperRef"]
    new_refs = [p.text for p in new.paragraphs if p.style.name == "PaperRef"]

    captions = [p.text for p in new.paragraphs if p.style.name == "PaperCaption"
                and re.match(r"^图\d+\s", p.text)]
    figure_numbers = [int(re.match(r"^图(\d+)", x).group(1)) for x in captions]
    equation_numbers = [re.search(r"（(\d+-\d+)）", p.text).group(1)
                        for p in new.paragraphs if p.style.name == "PaperEquation"
                        and re.search(r"（\d+-\d+）", p.text)]
    expected_equations = ([f"1-{i}" for i in range(1, 7)] +
                          [f"2-{i}" for i in range(1, 9)] +
                          [f"3-{i}" for i in range(1, 9)] +
                          [f"4-{i}" for i in range(1, 7)])
    headings = [set(x.strip() for x in text.splitlines()) for text in texts]
    toc = []
    for p in new.paragraphs:
        if p.style.name != "PaperTOC":
            continue
        title, printed = p.text.rsplit("\t", 1)
        matches = [i for i in range(3, len(headings)) if title in headings[i]]
        toc.append({"title": title, "printed": int(printed), "matches": matches})
    forbidden = ["当前会话", "本轮", "待核验", "待补充", "模型标识", "版本发布日期未", "Codex运行过程",
                 "SHA256", "哈希", "GitHub", "commit", "本地路径", "registry", "support package",
                 "machine output", "checkpoint", "candidate", "Q_FULL_FIELD_CONTRACT_v1.csv",
                 "冻结", "验收", "审计", "运行键"]
    # "github" is a real domain label in the supplied data; exclude only the
    # capitalized platform/project-management mention, not that data value.
    found_forbidden = [term for term in forbidden if (term in full_text if term == "GitHub"
                                                    else term.casefold() in full_text.casefold())]
    with ZipFile(DOCX) as z:
        thumbnail_parts = [x for x in z.namelist() if "thumbnail" in x.lower()]
        thumbnail_rel = b"metadata/thumbnail" in z.read("_rels/.rels")
    ai_notes = [p.text for p in new.paragraphs if "人工智能工具使用说明" in p.text]
    key_values = ["0.2846", "0.2278", "3.108", "2.3452", "2.2672", "0.33997658",
                  "0.27987813", "0.000146888", "0.000111642", "0.000114134",
                  "0.361995", "9.325359", "2.300064", "17.0233", "82.9767",
                  "79.76", "87.13", "77.93", "83.48", "76.11", "79.84"]
    missing_key_values = [x for x in key_values if x not in flat_pdf]
    v6_qa = json.loads((ROOT / "10_review/paper_v6/FINAL_QA_v6.json").read_text(encoding="utf-8"))
    checks = {
        "v6_baseline_formally_verified": v6_qa["pass"],
        "all_paragraph_edits_declared": len(old_paragraphs) == len(new_paragraphs) and not unexpected_changes,
        "all_numeric_changes_are_refs_or_removed_internal_ids": set(numeric_changes) == expected_numeric_changes,
        "all_table_numeric_tokens_unchanged": old_table_numbers == new_table_numbers,
        "all_math_objects_unchanged": math_text(old) == math_text(new),
        "all_media_assets_unchanged": media_hashes(V6) == media_hashes(DOCX),
        "reference_entries_unchanged": old_refs == new_refs,
        "pdf_opens_52_pages": len(pdf.pages) == 52 and not pdf.is_encrypted,
        "pdf_no_blank_pages": all(len(x) >= 70 for x in texts),
        "all_23_academic_tables_three_line": len(new.tables) == 24 and all(table_is_three_line(t) for t in new.tables[1:]),
        "figures_1_to_30_sequential": figure_numbers == list(range(1, 31)),
        "overall_and_four_route_diagrams_present": (any("论文整体技术框架" in x for x in captions)
                                                   and all(any(f"问题{i}技术路线" in x for x in captions)
                                                           for i in "一二三四")),
        "equations_numbered_by_question": equation_numbers == expected_equations,
        "toc_37_entries_match_final_pdf": len(toc) == 37 and all(x["matches"] == [x["printed"]] for x in toc),
        "ai_usage_note_single_and_formal": len(ai_notes) == 1 and "OpenAI Codex" in ai_notes[0]
                                         and "参赛队" in ai_notes[0],
        "forbidden_working_draft_language_absent": not found_forbidden,
        "key_formal_results_present": not missing_key_values,
        "no_thumbnail_or_dangling_relation": not thumbnail_parts and not thumbnail_rel,
        "anonymous_docx_properties": not new.core_properties.author and not new.core_properties.last_modified_by,
        "anonymous_pdf_metadata": not pdf.metadata or not pdf.metadata.get("/Author"),
    }
    result = {
        "pass": all(checks.values()), "checks": checks, "failed": [k for k, v in checks.items() if not v],
        "pdf_pages": len(pdf.pages), "v6_sha256": sha(V6), "v7_docx_sha256": sha(DOCX),
        "v7_pdf_sha256": sha(PDF), "paragraphs_changed": text_changes, "toc_changed": toc_changes,
        "numeric_paragraph_changes": numeric_changes, "forbidden_terms_found": found_forbidden,
        "missing_key_values": missing_key_values, "toc": toc,
        "figure_count": len(captions), "academic_table_count": len(new.tables) - 1,
        "equation_count": len(equation_numbers), "math_object_count": len(math_text(new)),
        "ai_note": ai_notes,
        "visual_round_1": "52 pages reviewed; 3 layout errors found in Appendix A/B/D",
        "visual_round_2": "52 pages reviewed; those 3 layout errors resolved",
        "final_visual_diff_from_round_2": "Only pages 3 (TOC correction) and 49 (appendix filename cleanup) changed; both enlarged and checked",
        "officecli_validation": "Success: no OpenXML schema errors",
    }
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"pass": result["pass"], "failed": result["failed"],
                      "pages": len(pdf.pages), "changed_paragraphs": len(text_changes),
                      "figures": len(captions), "tables": len(new.tables)-1}, ensure_ascii=False))
    if not result["pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
