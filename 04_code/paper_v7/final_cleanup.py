"""Publication cleanup of the frozen v6 DOCX; no model/result edits."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "11_delivery/v6/F_final_candidate_v6_获奖范文格式语言优化版.docx"


# Keys are stable v6 paragraph indices; the source excerpt is asserted before
# editing so a changed baseline cannot silently receive the wrong replacement.
REWRITES: dict[int, tuple[str, str]] = {
    73: ("数值一致性审计", "假设一：B1内部给定的N、D与Loss可用于同表关系估计。附件网格、完整轨迹及数值一致性检查支持这一用途；由于外部运行资料和共同Loss量尺尚不充分，结果用于附件内部推断。"),
    83: ("输入哈希、运行键", "表2按分析粒度组织附件。质量表按文本ID去重后比较扩展样本；配方表按运行索引匹配；B1保留完整训练轨迹，并以轨迹为单位划分验证集；能力表核对模型ID、参数、日期与任务量尺。各结果均报告有效样本及来源范围。"),
    134: ("自己的冻结参数", "质量标注表与训练配方表缺少同次训练的质量配对，因此分别建立质量描述与配比预测两条证据链。配比研究的背景见RegMix[7]；本文用题目附件的运行数据估计模型参数。A4/A5按索引匹配得到512组训练配方。p为17维非负配比，分量和为1；成分数据分析强调这一闭合约束[2]。考虑到零配比较多，采用欧氏单纯形正交对比，而不使用需额外处理零值的对数比坐标；该坐标不称为ILR。"),
    149: ("验收规则", "训练阶段五折综合分数为M0 1.001578、M1 0.750417、M2 0.749715。M2惩罚参数、模型结构和比较规则均在读取A6—A11的Loss前确定，M3预设残差升级条件未触发。A7按既定规则用于拟合外比较，不参与参数回填；因此它不是模型选择之后另设的最终测试集。"),
    156: ("冻结 M1 未重拟合", "经验支持由训练配比凸包、最近邻距离和训练分位阈值判定。1M的256组配比中，2组位于训练凸包内、252组处于预设近邻范围、2组在域外。A12配比与A13估算Loss、A14配比与A15估算Loss均按index成对匹配，分别包含10B/70B的63组运行。使用训练阶段确定的M1模型比较中心化响应形状、排序和几何支持，结果见表6；域外配比数量较少，结论以已检验范围为限。"),
    164: ("冻结 M1", "图7比较训练阶段确定的M1与M0的跨规模中心化误差比。1M呈现局部留出收益，60M部分迁移，1B未通过迁移检验；10B和70B估算结果用于压力检查。"),
    169: ("运行键", "本问把同来源曲面估计与跨资料质量、配比情景分开。B1八条训练轨迹同时提供N、D与同量尺Loss，因此可识别来源内衰减形状；A表提供配方及本来源Loss，B7提供半合成N—D—Q表。三者缺少同次训练实验及共同验证损失。五参数律因此先在B1中独立估计，质量和配比系数留作显式条件情景，以免把来源偏移解释为训练效应。"),
    247: ("本地冻结参数", "代入表7参数，内点预算弹性分别为0.451522和0.548478。达到D上界后，主要由N吸收新增预算；两上界同时活跃时预算变松。解析路径与51档预算、每档四个初值的数值结果逐点核对，最大目标差4.86×10⁻¹³、投影KKT残差2.78×10⁻¹⁷。两类求解一致性说明条件优化被准确实现；成本代理的现实解释仍取决于题设假设。"),
    278: ("冻结 51 档", "图17给出51档预算路径中基础训练、注意力和未使用预算的份额，并标示两个统计支持界的转移点。高预算的未用份额由支持界产生，不代表现实算力没有收益。"),
    324: ("用于匹配和审计", "日期采用评测文件时间，不以提交日或发布日期替代。C1中81个重复ID对应的162行整组排除；C8的1958个文件中4个截断文件保留记录但不解析，每个模型目录取最新可解析记录。C2只提供辅助日期；C3较早历史分数与六任务量尺不一致。C4中的数据、算力和开放权重字段用于资料匹配，可用训练D稀少且口径混合，因此不插补全样本训练量。"),
    339: ("冻结的 C1", "两类桥接在留出规模检验中均未优于均值基准，端部规模的表现也不稳定。七个高可比名称可在C8中找到，但符合预先规定的C1参数一致性条件者只有两条；样本又只覆盖一个家族，缺少留一家族及配对时间检验。因此当前结果不足以建立B1 Loss到C8 raw6分数的可靠映射。"),
    373: ("重新估计冻结模型", "第一类不确定性是给定动态模型与未来扰动时的条件模型区间。采用四周循环残差块重抽样[4]：1000次重抽后重新估计所选模型，并按残差标准差乘sqrt(1+h/4)的设定加入未来扰动。所得95%分位反映这一模型与扩散规则的条件变化；长期覆盖仍需同期限历史数据验证。"),
    391: ("冻结参数关联系数", "两期限参考中心分别为79.76、87.13；中度放缓为77.93、83.48；强放缓为76.11、79.84。情景差异随假定增长时长增加，来源于既定参数关联系数作用于外生路径。若训练D、算法或评测任务改变，能力路径亦会变化；这些数值应按明示条件解释。"),
    404: ("冻结规格", "不同敏感性因素作用于不同环节：质量权重和评价用途影响领域解释，成本函数与支持域影响资源配置，模型家族和预测结构影响远期投影。表20汇总实际比较的变化范围及其输出。“高低敏感”仅指设定的附件样本与模型规格，不能把机制差异并入单一抽样置信区间。"),
    438: ("冻结预测", "表B1列出1M留出运行全部13个验证领域的原尺度RMSE，与图6和表4使用同一训练阶段确定的模型预测。总体改善与领域间误差差异同时呈现。"),
    435: ("Q_FULL_FIELD_CONTRACT_v1.csv", "表A1分别列出全部22项字段族的原始解释角色、进入Q_full的规则和使用限制。逐字段目标与权重见随文提供的质量字段说明表。22项均进入W0等权综合评分；但原始值并非都具有越大越好的含义，转换后的效用也不能直接充当已识别的训练质量效应。19条不完整记录以中性0.5暂补并报告覆盖数；列表原分量保留，域级分布和有效样本数在随文附表提供。"),
    441: ("冻结参数", "配比模型以训练配比的正交坐标拟合各响应域，参数确定后在留出样本上检验；跨规模只比较中心化响应形状，支持分类使用训练配比几何。质量指标的参考分布、方向与权重不依赖留出Loss。"),
    442: ("保存验证索引", "标度律采用按轨迹等权的有界非线性最小二乘，并以完整轨迹重抽样评估条件稳定性。资源配置基线由式（3-5）计算；质量情景在固定Q时消去D，再比较端点与各局部谷。"),
    443: ("以哈希关联", "能力前沿由模型身份及六项原始任务分数构造；参数关联分解沿用同一分位插值权重。动态模型以既有历史端点滚动检验，条件区间与模型范围分别计算。计算程序与数据处理说明随文提供。"),
    444: ("冻结产物", "计算过程保留全精度，正文数值按所示位数四舍五入；因此用印刷参数回代时，末位差异不代表模型变化。"),
    445: ("SHA256", "主要数据处理方法、模型求解程序和完整结果随文提供，以便复核。"),
    446: ("当前会话模型标识", "人工智能工具使用说明：本文使用OpenAI Codex（OpenAI）辅助论文文字整理、排版、图表编排及程序编写与结果一致性核对。模型、公式、数值和结论由参赛队理解、复核并最终确定。"),
    448: ("项目支撑材料", "附录D汇总训练样本中的补充诊断，以说明正文采用多维质量表示和分领域响应的依据。各项数值由题目附件计算，未用于调整留出集模型；其解释范围与正文一致。"),
}


TABLE_CELL_REWRITES = {
    (2, 7, 1): "身份、筛选、日期及历史记录核对",
    (3, 0, 1): "候选与依据",
    (5, 2, 1): "不用于绝对预测评价",
    (5, 2, 2): "不用于绝对预测评价",
    (5, 3, 1): "不用于绝对预测评价",
    (5, 3, 2): "不用于绝对预测评价",
    (19, 0, 2): "主要结果",
    (19, 3, 2): "宏RMSE：\n0.000146888\n0.000111642\n0.000114134",
    (20, 5, 2): "低预算Loss：\n2.998935 →\n3.367161；参数带很窄",
    (23, 0, 1): "补充结果",
}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def numeric_tokens(value: str) -> Counter:
    return Counter(re.findall(r"(?<![A-Za-z])\d+(?:\.\d+)?(?:%)?", value))


def set_cell_font(cell, size: float) -> None:
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            rpr = run._r.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)
            rfonts.set(qn("w:eastAsia"), "宋体")


def set_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for ci, width in enumerate(widths):
        table.columns[ci].width = Inches(width)
        for row in table.rows:
            row.cells[ci].width = Inches(width)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(SOURCE)
    old_paras = list(doc.paragraphs)
    before_numeric = {}
    for index, (anchor, replacement) in REWRITES.items():
        paragraph = old_paras[index]
        if anchor not in paragraph.text:
            raise ValueError(f"Changed v6 baseline at paragraph {index}: {paragraph.text[:100]!r}")
        before_numeric[index] = numeric_tokens(paragraph.text)
        paragraph.text = replacement
        if index == 446:
            paragraph.clear()
            paragraph.add_run("人工智能工具使用说明：").bold = True
            paragraph.add_run(replacement.split("：", 1)[1])
    before_tables = [[cell.text for row in table.rows for cell in row.cells]
                     for table in doc.tables]
    for (ti, ri, ci), value in TABLE_CELL_REWRITES.items():
        doc.tables[ti].rows[ri].cells[ci].text = value

    for ti in (19, 20):
        set_widths(doc.tables[ti], [1.18, 1.15, 1.62, 1.15, 1.40])
        for row in doc.tables[ti].rows:
            for cell in row.cells:
                set_cell_font(cell, 9.5)
    # Long field identifiers in Appendix A need a little more room within the
    # existing four-column, three-line design; no field or numeric value changes.
    for row in doc.tables[21].rows[1:]:
        set_cell_font(row.cells[0], 9.5)
    # Keep Appendix A's heading, introduction, caption and first table row
    # together; keep Appendix D's compact table on its own final page.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(("附录 A 完整质量信号处理", "表A1分别列出")):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.startswith("附录 D 补充稳健性与探索性诊断"):
            paragraph.paragraph_format.page_break_before = True
    for cell in doc.tables[21].rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    for row in doc.tables[23].rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True

    doc.core_properties.title = "算力约束下大语言模型的质量评价、资源配置与能力前沿预测"
    for field in ("author", "last_modified_by", "comments", "keywords", "identifier"):
        setattr(doc.core_properties, field, "")
    doc.save(output)

    after_tables = [[cell.text for row in table.rows for cell in row.cells]
                    for table in doc.tables]
    table_numeric_equal = numeric_tokens("\n".join("\n".join(x) for x in before_tables)) == numeric_tokens(
        "\n".join("\n".join(x) for x in after_tables))
    if not table_numeric_equal:
        raise ValueError("Table numeric tokens changed")
    numeric_differences = {
        index: {"before": dict(before_numeric[index]),
                "after": dict(numeric_tokens(old_paras[index].text))}
        for index in REWRITES if before_numeric[index] != numeric_tokens(old_paras[index].text)
    }
    manifest = {
        "source": str(SOURCE.relative_to(ROOT)), "source_sha256": sha(SOURCE),
        "output": str(output), "output_sha256": sha(output),
        "paragraphs_edited": sorted(REWRITES), "table_cells_edited": [list(x) for x in TABLE_CELL_REWRITES],
        "numeric_differences_in_edited_paragraphs": numeric_differences,
        "table_numeric_tokens_preserved": table_numeric_equal,
    }
    (output.parent / "V7_EDIT_MANIFEST.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"docx": str(output), "paragraphs": len(REWRITES),
                      "table_cells": len(TABLE_CELL_REWRITES),
                      "numeric_paragraph_differences": sorted(numeric_differences),
                      "table_numeric_equal": table_numeric_equal}, ensure_ascii=False))


if __name__ == "__main__":
    main()
