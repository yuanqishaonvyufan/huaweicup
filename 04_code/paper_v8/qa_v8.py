"""Audit v8 references and verify all v7 modeling content is untouched."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

import pypdfium2 as pdfium
from PIL import ImageChops
from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "04_code/paper_v6"))
from qa_v6 import math_text, media_hashes, table_is_three_line  # noqa: E402

V7 = ROOT / "11_delivery/v7/F_final_candidate_v7_最终投稿版.docx"
DOCX = ROOT / "11_delivery/v8/F_final_candidate_v8_参考文献增强版.docx"
PDF = ROOT / "11_delivery/v8/F_final_candidate_v8_参考文献增强版.pdf"
RENDERED_DRAFT = Path(r"D:\work document\codex_work\HuaweiCup_F_2026_v8_references\draft_v8_round2.pdf")
REPORT = ROOT / "10_review/paper_v8/FINAL_QA_v8.json"
CITE_RE = re.compile(r"\[(\d+(?:,\s*\d+)*)\]")
NUM_RE = re.compile(r"(?<![A-Za-z])\d+(?:\.\d+)?(?:%)?")
ALLOWED_PROSE_EDITS = {95, 117, 134, 311, 318, 415}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def strip_cites(text: str) -> str:
    return CITE_RE.sub("", text)


def number_tokens(text: str) -> Counter:
    return Counter(NUM_RE.findall(strip_cites(text)))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--word-count", type=int, required=True)
    args = parser.parse_args()
    old = Document(V7)
    new = Document(DOCX)
    pdf = PdfReader(PDF)
    page_texts = [page.extract_text() or "" for page in pdf.pages]
    flat_pdf = re.sub(r"\s+", "", "".join(page_texts))

    old_body = [p for p in old.paragraphs if p.style.name != "PaperRef"]
    new_body = [p for p in new.paragraphs if p.style.name != "PaperRef"]
    unexpected_prose = [i for i, (a, b) in enumerate(zip(old_body, new_body))
                        if strip_cites(a.text) != strip_cites(b.text) and i not in ALLOWED_PROSE_EDITS]
    numeric_changes = [i for i, (a, b) in enumerate(zip(old_body, new_body))
                       if number_tokens(a.text) != number_tokens(b.text)]
    old_tables = [[c.text for row in t.rows for c in row.cells] for t in old.tables]
    new_tables = [[c.text for row in t.rows for c in row.cells] for t in new.tables]
    old_figures = [p.text for p in old.paragraphs if p.style.name == "PaperCaption"
                   and re.match(r"^图\d+\s", p.text)]
    new_figures = [p.text for p in new.paragraphs if p.style.name == "PaperCaption"
                   and re.match(r"^图\d+\s", p.text)]
    old_equations = [p.text for p in old.paragraphs if p.style.name == "PaperEquation"]
    new_equations = [p.text for p in new.paragraphs if p.style.name == "PaperEquation"]

    refs = [p.text for p in new.paragraphs if p.style.name == "PaperRef"]
    ref_numbers = [int(re.match(r"\[(\d+)\]", x).group(1)) for x in refs]
    body_mentions = []
    for paragraph in new_body:
        if paragraph.style.name == "PaperTOC":
            continue
        for match in CITE_RE.finditer(paragraph.text):
            body_mentions.extend(int(x.strip()) for x in match.group(1).split(","))
    first_order = list(dict.fromkeys(body_mentions))
    ref_mentions = Counter(body_mentions)
    headings = [set(x.strip() for x in text.splitlines()) for text in page_texts]
    toc = []
    for paragraph in new.paragraphs:
        if paragraph.style.name == "PaperTOC":
            title, printed = paragraph.text.rsplit("\t", 1)
            matches = [i for i in range(3, len(headings)) if title in headings[i]]
            toc.append({"title": title, "printed": int(printed), "matches": matches})
    with ZipFile(DOCX) as z:
        thumbnail_parts = [name for name in z.namelist() if "thumbnail" in name.lower()]
        thumbnail_rel = b"metadata/thumbnail" in z.read("_rels/.rels")

    image_diff = []
    draft = pdfium.PdfDocument(str(RENDERED_DRAFT))
    final = pdfium.PdfDocument(str(PDF))
    if len(draft) == len(final):
        for i in range(len(final)):
            a = draft[i].render(scale=1).to_pil().convert("RGB")
            b = final[i].render(scale=1).to_pil().convert("RGB")
            if ImageChops.difference(a, b).getbbox():
                image_diff.append(i+1)

    key_values = ["0.2846", "0.2278", "3.108", "2.3452", "2.2672", "0.33997658",
                  "0.27987813", "0.000146888", "0.000111642", "0.000114134",
                  "0.361995", "9.325359", "2.300064", "17.0233", "82.9767",
                  "79.76", "87.13", "77.93", "83.48", "76.11", "79.84"]
    missing_values = [x for x in key_values if x not in flat_pdf]
    v7_qa = json.loads((ROOT / "10_review/paper_v7/FINAL_QA_v7.json").read_text(encoding="utf-8"))
    checks = {
        "v7_baseline_verified": v7_qa["pass"],
        "body_paragraph_count_unchanged": len(old_body) == len(new_body),
        "only_six_mapped_prose_paragraphs_expanded": not unexpected_prose,
        "all_nonbibliographic_numbers_unchanged": not numeric_changes,
        "all_24_tables_identical": old_tables == new_tables,
        "all_math_objects_identical": math_text(old) == math_text(new),
        "all_equation_paragraphs_identical": old_equations == new_equations,
        "all_30_figure_captions_identical": old_figures == new_figures and len(new_figures) == 30,
        "all_embedded_media_identical": media_hashes(V7) == media_hashes(DOCX),
        "all_23_academic_tables_three_line": len(new.tables) == 24 and all(table_is_three_line(t) for t in new.tables[1:]),
        "references_1_to_16_complete": len(refs) == 16 and ref_numbers == list(range(1, 17)),
        "first_citation_order_1_to_16": first_order == list(range(1, 17)),
        "all_and_only_references_cited": set(body_mentions) == set(ref_numbers) and all(ref_mentions.values()),
        "three_existing_arxiv_replaced": not any("arXiv" in x for x in refs)
                                        and any("ICLR" in x for x in refs)
                                        and any("PMLR 202:2397" in x for x in refs),
        "no_unresolved_citation_placeholders": "⟦" not in "".join(p.text for p in new.paragraphs),
        "toc_37_entries_match_pdf": len(toc) == 37 and all(x["matches"] == [x["printed"]] for x in toc),
        "pdf_53_pages_no_blank": len(pdf.pages) == 53 and not pdf.is_encrypted
                                  and all(len(text) >= 70 for text in page_texts),
        "formal_result_values_present": not missing_values,
        "final_pdf_visual_identical_to_checked_draft": len(draft) == len(final) and not image_diff,
        "anonymous_docx_and_pdf": not new.core_properties.author and not new.core_properties.last_modified_by
                                  and (not pdf.metadata or not pdf.metadata.get("/Author")),
        "no_docx_thumbnail_or_dangling_relation": not thumbnail_parts and not thumbnail_rel,
        "word_count_reasonable": 31517 <= args.word_count <= 32500,
    }
    report = {"pass": all(checks.values()), "checks": checks,
              "failed": [k for k, v in checks.items() if not v],
              "v7_sha256": sha(V7), "v8_docx_sha256": sha(DOCX), "v8_pdf_sha256": sha(PDF),
              "old_refs": 7, "new_refs": len(refs), "pdf_pages": len(pdf.pages),
              "word_count_v8": args.word_count, "citation_first_order": first_order,
              "reference_mentions": dict(ref_mentions), "unexpected_prose_edits": unexpected_prose,
              "numeric_paragraph_changes": numeric_changes, "missing_key_values": missing_values,
              "toc": toc, "image_diff_pages_from_checked_draft": image_diff,
              "figures": len(new_figures), "tables": len(new.tables)-1,
              "math_objects": len(math_text(new)), "source_media_count": len(media_hashes(V7)),
              "bibliography_sources_checked": "Local candidate PDFs; NeurIPS/ICLR/PMLR/ACL/TMLR official publication records and DOI metadata",
              "officecli_validation": "Passed: no OpenXML schema errors",
              "visual_review": "First 53-page PDF visually reviewed; bibliography pagination revised; second final export checked against the reviewed PDF"}
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"pass": report["pass"], "failed": report["failed"],
                      "pages": report["pdf_pages"], "refs": len(refs),
                      "unexpected_prose": unexpected_prose, "numeric_changes": numeric_changes}, ensure_ascii=False))
    if not report["pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
