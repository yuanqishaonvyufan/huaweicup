"""Insert source-matched citations into v7 without changing its models/results."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "11_delivery/v7/F_final_candidate_v7_最终投稿版.docx"
OLD_KEYS = {1: "Hoffmann", 2: "Aitchison", 3: "Pythia", 4: "Kunsch",
            5: "Problem", 6: "DataCard", 7: "RegMix"}

REFERENCES = {
    "Hoffmann": "Hoffmann J, Borgeaud S, Mensch A, et al. An Empirical Analysis of Compute-Optimal Large Language Model Training. Advances in Neural Information Processing Systems, 35:30016–30030, 2022. DOI: 10.52202/068431-2176.",
    "DataComp": "Li J, Fang A, Smyrnis G, et al. DataComp-LM: In Search of the Next Generation of Training Sets for Language Models. Advances in Neural Information Processing Systems, 37:14200–14282, 2024. DOI: 10.52202/079017-0455.",
    "RefinedWeb": "Penedo G, Malartic Q, Hesslow D, et al. The RefinedWeb Dataset for Falcon LLM: Outperforming Curated Corpora with Web Data Only. Advances in Neural Information Processing Systems, 36:79155–79172, 2023. DOI: 10.52202/075280-3464.",
    "FineWeb": "Penedo G, Kydlíček H, Ben Allal L, et al. The FineWeb Datasets: Decanting the Web for the Finest Text Data at Scale. Advances in Neural Information Processing Systems, 37:30811–30849, 2024. DOI: 10.52202/079017-0970.",
    "Problem": "中国研究生数学建模竞赛组委会. 2026年F题题面及配套数据说明[Z]. 竞赛附件, 2026.",
    "DataCard": "opendatalab. SlimPajama Meta-rater 数据卡[DB/OL]. Hugging Face, 访问日期2026-09-25. https://huggingface.co/datasets/opendatalab/SlimPajama-Meta-rater/blob/main/README.md.",
    "Lee": "Lee K, Ippolito D, Nystrom A, et al. Deduplicating Training Data Makes Language Models Better. Proceedings of the 60th Annual Meeting of the Association for Computational Linguistics, 8424–8445, 2022. DOI: 10.18653/v1/2022.acl-long.577.",
    "DSIR": "Xie S M, Santurkar S, Ma T, et al. Data Selection for Language Models via Importance Resampling. Advances in Neural Information Processing Systems, 36:34201–34227, 2023. DOI: 10.52202/075280-1482.",
    "DoReMi": "Xie S M, Pham H, Dong X, et al. DoReMi: Optimizing Data Mixtures Speeds Up Language Model Pretraining. Advances in Neural Information Processing Systems, 36:69798–69818, 2023. DOI: 10.52202/075280-3059.",
    "RegMix": "Liu Q, Zheng X, Muennighoff N, et al. RegMix: Data Mixture as Regression for Language Model Pre-training. International Conference on Learning Representations (ICLR), 2025.",
    "Aitchison": "Aitchison J. The Statistical Analysis of Compositional Data. Journal of the Royal Statistical Society: Series B (Methodological), 44(2):139–177, 1982. DOI: 10.1111/j.2517-6161.1982.tb01195.x.",
    "Pythia": "Biderman S, Schoelkopf H, Anthony Q G, et al. Pythia: A Suite for Analyzing Large Language Models Across Training and Scaling. Proceedings of the 40th International Conference on Machine Learning, PMLR 202:2397–2430, 2023. https://proceedings.mlr.press/v202/biderman23a.html.",
    "HELM": "Liang P, Bommasani R, Lee T, et al. Holistic Evaluation of Language Models. Transactions on Machine Learning Research, 2023. https://openreview.net/forum?id=iO4LZibEqW.",
    "BIG": "Srivastava A, Rastogi A, Rao A, et al. Beyond the Imitation Game: Quantifying and Extrapolating the Capabilities of Language Models. Transactions on Machine Learning Research, 2023. https://openreview.net/forum?id=uyTL5Bvosj.",
    "Kunsch": "Künsch H R. The Jackknife and the Bootstrap for General Stationary Observations. The Annals of Statistics, 17(3):1217–1241, 1989. DOI: 10.1214/aos/1176347265.",
    "Deng": "Deng C, Zhao Y, Tang X, et al. Investigating Data Contamination in Modern Benchmarks for Large Language Models. Proceedings of the 2024 Conference of the North American Chapter of the Association for Computational Linguistics: Human Language Technologies, 8706–8719, 2024. DOI: 10.18653/v1/2024.naacl-long.482.",
}

CITE_PATTERN = re.compile(r"⟦([A-Za-z,]+)⟧|\[(\d+)\]")


def replace_run(paragraph, old: str, new: str) -> None:
    hits = [run for run in paragraph.runs if old in run.text]
    if len(hits) != 1 or hits[0].text.count(old) != 1:
        raise ValueError(f"Anchor {old!r} has {len(hits)} run matches: {paragraph.text[:90]}")
    hits[0].text = hits[0].text.replace(old, new, 1)


def citation_keys(text: str):
    for match in CITE_PATTERN.finditer(text):
        if match.group(1):
            yield from match.group(1).split(",")
        else:
            yield OLD_KEYS[int(match.group(2))]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(BASE)
    p = doc.paragraphs

    # Remove one redundant foundational citation. The same NeurIPS paper stays
    # cited at the introduction and at the actual Q2 model statement.
    replace_run(p[64], "[1]", "")
    replace_run(p[95], "。先按竞赛附件",
                "。受控语料研究将过滤、去重和混合设为可评估的数据设计环节⟦DataComp⟧；RefinedWeb与FineWeb-Edu分别提供网页语料整理和教育文本筛选的研究案例⟦RefinedWeb,FineWeb⟧。先按竞赛附件")
    replace_run(p[95], "。列表型指标",
                "。近重复语料可能影响训练与评测解释，去重作为语料整理步骤已有实证依据⟦Lee⟧。列表型指标")
    replace_run(p[117], "不视为三份独立质量证据。",
                "不视为三份独立质量证据。DSIR原方法以目标分布匹配进行数据选择⟦DSIR⟧，其用途与通用质量排序不同。")
    replace_run(p[134], "配比研究的背景见",
                "DoReMi与RegMix分别以领域重加权和代理回归研究训练数据配比")
    replace_run(p[134], "RegMix[7]", "⟦DoReMi,RegMix⟧")
    replace_run(p[311], "为预测目标。首先用",
                "为预测目标。多任务评测研究强调明确场景与指标口径⟦HELM⟧。首先用")
    replace_run(p[318], "总分上升并不表示每项能力以相同比例改善。",
                "总分上升并不表示每项能力以相同比例改善。BIG-bench的跨任务研究说明不同任务的规模响应可能呈现差异⟦BIG⟧。")
    replace_run(p[415], "模型得分上界当作真实技术饱和证据。",
                "模型得分上界当作真实技术饱和证据。基准污染也可能影响得分解释⟦Deng⟧；本文未对C8与训练语料的重叠作独立检验。")

    body = [paragraph for paragraph in doc.paragraphs if paragraph.style.name != "PaperRef"]
    first_order = []
    placement = {}
    for index, paragraph in enumerate(body):
        if paragraph.style.name == "PaperTOC":
            continue
        for key in citation_keys(paragraph.text):
            if key not in REFERENCES:
                raise ValueError(f"No bibliography entry for {key}")
            if key not in first_order:
                first_order.append(key)
            placement.setdefault(key, []).append(index)
    if set(first_order) != set(REFERENCES):
        raise ValueError(f"Uncited bibliography entries: {set(REFERENCES)-set(first_order)}")
    numbers = {key: index+1 for index, key in enumerate(first_order)}

    def replace_cite(match: re.Match) -> str:
        keys = match.group(1).split(",") if match.group(1) else [OLD_KEYS[int(match.group(2))]]
        return "[" + ", ".join(str(numbers[key]) for key in keys) + "]"

    for paragraph in body:
        for run in paragraph.runs:
            if CITE_PATTERN.search(run.text):
                run.text = CITE_PATTERN.sub(replace_cite, run.text)

    reference_paragraphs = [paragraph for paragraph in doc.paragraphs
                            if paragraph.style.name == "PaperRef"]
    if len(reference_paragraphs) != 7:
        raise ValueError(f"Expected 7 v7 references, found {len(reference_paragraphs)}")
    anchor = reference_paragraphs[-1]
    for index, key in enumerate(first_order):
        text = f"[{index+1}] {REFERENCES[key]}"
        if index < len(reference_paragraphs):
            reference_paragraphs[index].text = text
        else:
            paragraph = doc.add_paragraph(style="PaperRef")
            paragraph.add_run(text)
            anchor._p.addnext(paragraph._p)
            anchor = paragraph

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "PaperRef":
            paragraph.paragraph_format.keep_together = True
        elif paragraph.text.startswith("附录 A 完整质量信号处理"):
            paragraph.paragraph_format.page_break_before = True

    doc.core_properties.title = "算力约束下大语言模型的质量评价、资源配置与能力前沿预测"
    for field in ("author", "last_modified_by", "comments", "keywords", "identifier"):
        setattr(doc.core_properties, field, "")
    doc.save(output)
    manifest = {"source": str(BASE), "output": str(output), "reference_count_v7": 7,
                "reference_count_v8": len(first_order), "ordered_keys": first_order,
                "first_appearance_paragraphs": {key: locs[0] for key, locs in placement.items()},
                "citation_paragraphs": placement,
                "body_insertions": [95, 117, 134, 311, 318, 415],
                "redundant_citation_removed": 64}
    (output.parent / "V8_BUILD_MANIFEST.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"references": len(first_order), "order": first_order}, ensure_ascii=False))


if __name__ == "__main__":
    main()
