"""Build the F-problem v7 source listings from frozen, tracked Python files.

This is a packaging utility. It never executes the modeling scripts or rewrites
their inputs, outputs, or source files.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
BASELINE = "43d01e2"  # v7 submission paper commit; source paths below are unchanged.

# In execution/dependency order. The original paths are the executable sources;
# the Markdown and Word listings are byte-traceable copies for review.
SOURCES = {
    "Q1": [
        ("04_code/modeling_phase1/round3_prepare_inputs.py", "核对并准备 A 类原始输入，形成描述分析使用的规范输入。"),
        ("04_code/modeling_phase1/round3_freeze_signal_set.py", "固定质量信号集合、字段语义和后续描述模型的输入口径。"),
        ("04_code/modeling_phase1/round3_q1_descriptive.py", "计算五维质量画像和 DQ0 等描述性对照。"),
        ("04_code/modeling_phase1/round4_quality_closure.py", "核对 22 项信号角色及领域层质量描述，支持最终评分的解释边界。"),
        ("04_code/modeling_phase1/round4_train_p_response.py", "在 A4/A5 上训练 17 域配比到 Loss 的 M0/M1 候选；验证集不参与拟合。"),
        ("04_code/modeling_phase1/round4_validate_p_response.py", "在 A6—A11 上验证冻结配比模型，并检查跨规模迁移和经验支持。"),
        ("04_code/paper_round8/q1_task_repair.py", "生成 22 项用途规则评分 Q_full，并对 A12—A15 作估算外推压力检验。"),
        ("04_code/modeling_phase1/round4_q1_figures.py", "从已登记结果绘制问题一配比响应和支持范围图。"),
        ("04_code/paper_v5/supplementary_q1_vis.py", "从冻结表绘制最终稿使用的质量排序与补充诊断图。"),
    ],
    "Q2": [
        ("04_code/modeling_phase2/round5_prepare.py", "准备 B 类标度律输入并记录来源与字段校验。"),
        ("04_code/modeling_phase2/round5_fit.py", "拟合五参数 N–D 加性幂律，执行轨迹分组验证与参数稳定性计算。"),
        ("04_code/modeling_phase2/round5_scenarios.py", "计算质量校准、跨来源形状诊断和条件替代情景。"),
        ("04_code/visualization/round5_figures.py", "绘制问题二标度律、验证、弹性和情景图。"),
        ("04_code/paper_v5/supplementary_q2_vis.py", "显示冻结的整轨迹重抽样参数分布。"),
    ],
    "Q3": [
        ("04_code/modeling_phase3/round6_core.py", "实现题设 FLOPs 成本、N–D 损失、解析最优、质量情景和 KKT 诊断。"),
        ("04_code/modeling_phase3/round6_baseline.py", "求解 51 个预算点的 N–D 基线最优路径。"),
        ("04_code/modeling_phase3/round6_scenarios.py", "计算质量、上下文长度、配比和供给约束的条件情景。"),
        ("04_code/modeling_phase3/round6_uncertainty.py", "传播联合参数向量的不确定性并计算情景包络。"),
        ("04_code/visualization/round6_figures.py", "绘制问题三预算路径、约束与不确定性图。"),
        ("04_code/paper_v5/supplementary_q3_vis.py", "显示最终稿使用的预算构成和活跃约束补充图。"),
    ],
    "Q4": [
        ("04_code/modeling_phase4/common.py", "提供问题四共享路径、数据读写和规范化工具。"),
        ("04_code/modeling_phase4/round7_mapping.py", "核对 C 类附件字段、模型实体和评测记录映射。"),
        ("04_code/modeling_phase4/round7_freeze.py", "构建供前沿、桥接和预测使用的冻结规范输入。"),
        ("04_code/modeling_phase4/round7_models.py", "计算评测前沿、描述性分解并检验 Loss–Benchmark 桥接。"),
        ("04_code/modeling_phase4/round7_forecast.py", "比较动态模型并给出 12/24 个月条件预测及区间。"),
        ("04_code/paper_round8/q4_task_repair.py", "计算可用训练 D 子样本诊断和外生正增长放缓情景。"),
        ("04_code/modeling_phase4/round7_figures.py", "绘制问题四前沿、桥接、分解、预测及敏感性图。"),
        ("04_code/paper_round8/revised_figures.py", "从冻结结果重导出最终稿的修订图。"),
        ("04_code/paper_v5/supplementary_q4_vis.py", "绘制最终稿的放缓情景与任务分量补充图。"),
    ],
}

EVIDENCE = {
    "Q1": [
        "03_models/modeling_phase1/q1/round4/P_RESPONSE_VALIDATION_METRICS_v2.json",
        "06_results/raw/Q1_TASK_REPAIR_20260925_v1/summary.json",
    ],
    "Q2": [
        "06_results/raw/EXP-Q2-ND-R5-20260924-v1/summary.json",
        "06_results/raw/SCEN-Q2-R5-20260924-v3/summary.json",
    ],
    "Q3": [
        "06_results/raw/EXP-Q3-BASE-R6-20260924-v1/summary.json",
        "06_results/raw/SCEN-Q3-R6-20260924-v1/summary.json",
        "06_results/raw/UNC-Q3-R6-20260924-v1/summary.json",
    ],
    "Q4": [
        "06_results/raw/EXP-Q4-R7-20260925-v1/bridge_summary.json",
        "06_results/raw/EXP-Q4-R7-20260925-v1/forecast_summary.json",
        "06_results/raw/Q4_TASK_REPAIR_20260925_v1/summary.json",
    ],
}

NAMES = {
    "Q1": "问题一：数据质量与配比响应",
    "Q2": "问题二：N–D 标度律与条件替代",
    "Q3": "问题三：算力预算下的资源配置",
    "Q4": "问题四：能力前沿、桥接与预测",
}


def source_record(section: str, number: int, path: str, purpose: str) -> dict:
    target = ROOT / path
    if not target.is_file():
        raise FileNotFoundError(target)
    data = target.read_bytes()
    text = data.decode("utf-8-sig")
    compile(text, str(target), "exec")
    return {
        "section": section,
        "number": number,
        "path": path,
        "purpose": purpose,
        "sha256": hashlib.sha256(data).hexdigest(),
        "lines": len(text.splitlines()),
        "bytes": len(data),
        "text": text,
    }


def records(section: str) -> list[dict]:
    return [source_record(section, i, path, purpose)
            for i, (path, purpose) in enumerate(SOURCES[section], 1)]


def prompt(section: str) -> str:
    return (HERE / "prompts" / f"{section}.md").read_text(encoding="utf-8").strip()


def render_markdown(section: str) -> None:
    rows = records(section)
    for path in EVIDENCE[section]:
        if not (ROOT / path).is_file():
            raise FileNotFoundError(ROOT / path)
    out = [f"# {NAMES[section]}：最终版源码", "",
           f"对应论文：`11_delivery/v7/F_final_candidate_v7_最终投稿版.docx`。源码基线：`{BASELINE}`。",
           "下列代码逐字取自项目原路径；SHA-256 按原始文件字节计算。冻结结果原件不在本页重算或覆盖。", "",
           "## 可直接粘贴到 Codex 的提示词", "", prompt(section), "",
           "## 正式结果核对入口", ""]
    out += [f"- `{path}`" for path in EVIDENCE[section]]
    out += [""]
    for row in rows:
        out += [f"## {section}-{row['number']:02d} {Path(row['path']).name}", "",
                f"用途：{row['purpose']}", "",
                f"源路径：`{row['path']}`  ",
                f"SHA-256：`{row['sha256']}`  ",
                f"行数：{row['lines']}", "", "`````python",
                row["text"].rstrip("\n"), "`````", ""]
    (HERE / f"{section}.md").write_text("\n".join(out), encoding="utf-8", newline="\n")
    manifest = [{k: v for k, v in row.items() if k != "text"} for row in rows]
    (HERE / f"{section}_manifest.json").write_text(
        json.dumps({"paper": "v7", "baseline": BASELINE, "section": section,
                    "sources": manifest, "evidence": EVIDENCE[section]},
                   ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


def east_asia_font(style, name: str) -> None:
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), name)


def build_docx(target: Path) -> None:
    all_rows = {section: records(section) for section in SOURCES}
    for section in SOURCES:
        for path in EVIDENCE[section]:
            if not (ROOT / path).is_file():
                raise FileNotFoundError(ROOT / path)
        prompt(section)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = Inches(.67)
    sec.bottom_margin = Inches(.64)
    sec.left_margin = Inches(.72)
    sec.right_margin = Inches(.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    east_asia_font(normal, "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    for name, size, before, after in (("Title", 20, 0, 13),
                                      ("Heading 1", 15, 16, 9),
                                      ("Heading 2", 11.5, 10, 5)):
        st = doc.styles[name]
        st.font.name = "Microsoft YaHei"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(25, 49, 77)
        east_asia_font(st, "Microsoft YaHei")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    code = doc.styles.add_style("CodeLine", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.3)
    code.font.color.rgb = RGBColor(32, 39, 48)
    east_asia_font(code, "Microsoft YaHei")
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.left_indent = Inches(.12)
    code.paragraph_format.widow_control = False
    pstyle = doc.styles.add_style("CodexPrompt", WD_STYLE_TYPE.PARAGRAPH)
    pstyle.base_style = normal
    pstyle.font.size = Pt(9.5)
    pstyle.paragraph_format.space_after = Pt(3)
    pstyle.paragraph_format.left_indent = Inches(.15)
    title = doc.add_paragraph(style="Title")
    title.add_run("2026 华为杯 F 题最终版代码汇编")
    doc.add_paragraph("对应 v7 最终投稿版 · 四问模型、验证与作图源码")
    total = sum(len(rows) for rows in all_rows.values())
    lines = sum(row["lines"] for rows in all_rows.values() for row in rows)
    doc.add_paragraph(f"源码基线：{BASELINE}    源文件：{total} 个    合计：{lines} 行")
    doc.add_paragraph("代码按原路径与原始字节 SHA-256 登记；Word 中保留可复制的完整源码。运行产生的历史结果已有固定编号与哈希，复算时应使用隔离副本。")
    doc.add_heading("目录与使用方法", 1)
    for section in SOURCES:
        doc.add_paragraph(f"{section}  {NAMES[section]}（{len(all_rows[section])} 个文件）")
    doc.add_paragraph("复制相应问题开头的提示词到 Codex，即可要求它核对源文件、输入与正式结果。直接复制 Word 中的代码时，须保持原脚本所在目录和相对路径；脚本的固定运行编号可能拒绝覆盖已有结果。")
    for section, rows in all_rows.items():
        doc.add_page_break()
        doc.add_heading(NAMES[section], 1)
        doc.add_paragraph("正式结果核对入口：")
        for path in EVIDENCE[section]:
            doc.add_paragraph(path, style="CodexPrompt")
        doc.add_heading("可直接粘贴到 Codex 的提示词", 2)
        for line in prompt(section).splitlines():
            doc.add_paragraph(line, style="CodexPrompt")
        for row in rows:
            doc.add_heading(f"{section}-{row['number']:02d}  {Path(row['path']).name}", 2)
            doc.add_paragraph(f"用途：{row['purpose']}")
            doc.add_paragraph(f"源路径：{row['path']}")
            doc.add_paragraph(f"SHA-256：{row['sha256']}    行数：{row['lines']}")
            for line in row["text"].splitlines():
                para = doc.add_paragraph(style="CodeLine")
                para.add_run(line)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(target)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--section", choices=list(SOURCES))
    parser.add_argument("--docx", type=Path)
    args = parser.parse_args()
    if not args.section and not args.docx:
        parser.error("Specify --section or --docx")
    if args.section:
        render_markdown(args.section)
    if args.docx:
        build_docx(args.docx)


if __name__ == "__main__":
    main()
