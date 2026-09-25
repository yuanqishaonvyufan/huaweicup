"""Build manuscript from frozen text/results; never fit a model.
AI-assisted code: Codex / GPT-6, OpenAI; release date independently unverified.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import re, json, csv, hashlib, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree
from latex2mathml.converter import convert
from PIL import Image

ROOT=Path(__file__).resolve().parents[2]
SRC=ROOT/'08_paper/round8'; OUT=ROOT/'11_delivery/round8'
OUT.mkdir(exist_ok=True,parents=True)
TEMPLATE=SRC/'template_reference.docx'
def jwrite(path,x): path.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def csvrows(p): return list(csv.DictReader((ROOT/p).open(encoding='utf-8-sig',newline='')))
source_files=['00_front.md','05_Q1.md','06_Q2.md','07_Q3.md','08_Q4.md','09_back.md']
raw='\n\n'.join((SRC/p).read_text(encoding='utf-8') for p in source_files)
names=['教育价值','英语流畅度','清洁度','可读性','推理复杂度','专业性','书籍相似性','百科相似性','数学相似性','四面向质量评级','无广告程度','词数','句数','一元熵','不重复词比例','非字母词比例','最高频二元组字符占比','最高频三元组字符占比','大写字母比例','行尾标点比例','数字字符比例','平均词长']
processing=['原评分；正向主坐标','正负logit差；正向主坐标','概率加权等级；正向主坐标','概率加权等级；正向主坐标','保留等级；依用途分析','保留等级；依用途分析','尺度未核；不进入总分','与书籍列近重复；单列诊断','与书籍列近重复；单列诊断','保留四分量；不强行平均','无广告减有广告logit；主坐标','log1p长度协变量','log1p长度协变量','原值/秩；基数未核','原值/秩；依语言与用途','原值/秩；代码域另析','原值保留；不裁剪超界','原值保留；不裁剪超界','原值/秩；用途敏感','原值/秩；用途敏感','原值/秩；用途敏感','字符/词；用途敏感']
roles=csvrows('03_models/modeling_phase1/q1/round4/Q1_FULL22_ROLE_MAP_v1.csv')
signals='表A1 完整22项信号的含义与使用方式\n\n|序号及字段族|含义|处理与限制|\n|---|---|---|\n'
for i,(r,n,p) in enumerate(zip(roles,names,processing),1): signals+=f'|{i}. {r["signal_id"]}|{n}|{p}|\n'
dom=csvrows('03_models/modeling_phase1/q1/round4/P_RESPONSE_DOMAIN_VALIDATION_v2.csv')
dom0={r['domain']:r for r in dom if r['scale']=='1M' and r['model']=='M0'}
domain_table='表B1 1M留出逐领域预测误差\n\n|验证领域|M0 RMSE|M1 RMSE|M1/M0|\n|---|---|---|---|\n'
for r in dom:
 if r['scale']=='1M' and r['model']=='M1':
  domain_table+=f'|{r["domain"]}|{float(dom0[r["domain"]]["RMSE_raw"]):.4f}|{float(r["RMSE_raw"]):.4f}|{float(r["RMSE_ratio_to_M0_raw"]):.4f}|\n'
raw=raw.replace('@TABLE_SIGNALS',signals).replace('@TABLE_DOMAINS',domain_table)
(OUT/'F_final_manuscript.md').write_text(raw,encoding='utf-8')
doc=Document(TEMPLATE)
cover_nodes=list(doc.element.body)[:7]
cover_xml=[etree.tostring(x) for x in cover_nodes]
for e in list(doc.element.body)[7:]:
 if e.tag!=qn('w:sectPr'): doc.element.body.remove(e)
# The retained template's cover is preserved; replace its blank abstract/body slots.
first=doc.sections[0]
for foot in [first.footer,first.first_page_footer,first.even_page_footer]:
 for e in list(foot._element): foot._element.remove(e)
for head in [first.header,first.first_page_header,first.even_page_header]:
 for e in list(head._element): head._element.remove(e)
section=doc.add_section(WD_SECTION_START.NEW_PAGE)
section.footer.is_linked_to_previous=False
section.header.is_linked_to_previous=False
section.different_first_page_header_footer=False
for old in section._sectPr.findall(qn('w:pgNumType')):section._sectPr.remove(old)
pg=OxmlElement('w:pgNumType'); pg.set(qn('w:start'),'1'); section._sectPr.append(pg)
for e in list(section.header._element):section.header._element.remove(e)
# Remove references entirely: LibreOffice can synthesize a bordered Header style
# even for an empty header part retained by the old binary template.
for sec in doc.sections:
 for e in list(sec._sectPr.findall(qn('w:headerReference'))):sec._sectPr.remove(e)
for rid,rel in list(doc.part.rels.items()):
 if rel.reltype.endswith('/header'):doc.part.drop_rel(rid)
p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
for c in doc.core_properties.__class__.__dict__:
 if c in ['author','last_modified_by','comments','keywords','subject','category','identifier']:
  setattr(doc.core_properties,c,'')
doc.core_properties.title='算力约束下大语言模型的质量评价资源配置与能力前沿预测'
doc.core_properties.revision=1
def style(name,size=12,bold=False,align=None):
 s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
 rpr=s.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
 rf.set(qn('w:eastAsia'),'黑体' if name in ['Title','PaperH1'] else '宋体')
 pf=s.paragraph_format;pf.line_spacing=1;pf.space_after=Pt(3);pf.space_before=Pt(0)
 if align is not None: pf.alignment=align
 return s
body=style('PaperBody');body.paragraph_format.first_line_indent=Pt(24)
style('Title',16,True,WD_ALIGN_PARAGRAPH.CENTER).paragraph_format.space_after=Pt(12)
h1=style('PaperH1',14,True,WD_ALIGN_PARAGRAPH.CENTER);h1.paragraph_format.space_before=Pt(12);h1.paragraph_format.space_after=Pt(8);h1.paragraph_format.keep_with_next=True
h2=style('PaperH2',12,True);h2.paragraph_format.space_before=Pt(8);h2.paragraph_format.space_after=Pt(5);h2.paragraph_format.keep_with_next=True
style('PaperCaption',12,False,WD_ALIGN_PARAGRAPH.CENTER)
style('PaperTable',12);style('PaperEquation',12,False,WD_ALIGN_PARAGRAPH.CENTER)
style('PaperRef',12)
width_cm=section.page_width.cm-section.left_margin.cm-section.right_margin.cm
def paragraph(text,sty='PaperBody'):
 p=doc.add_paragraph(text,style=sty)
 p.paragraph_format.widow_control=True
 return p
def table(data):
 t=doc.add_table(rows=1,cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 # More room for long descriptions and appendix field names.
 weights={3:[.32,.25,.43],4:[.31,.23,.23,.23],5:[.18,.17,.18,.18,.29]}.get(len(data[0]),[1/len(data[0])]*len(data[0]))
 if '字段族' in data[0][0]:weights=[.42,.20,.38]
 if '验证方式' in data[0][0]:weights=[.19,.16,.16,.29,.20]
 for c,w in zip(t.columns,weights):c.width=Cm(width_cm*w)
 for ri,row in enumerate(data):
  cells=t.rows[0].cells if ri==0 else t.add_row().cells
  pr=t.rows[ri]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
  if ri==0:rep=OxmlElement('w:tblHeader');pr.append(rep)
  for ci,(cell,value) in enumerate(zip(cells,row)):
   cell.width=Cm(width_cm*weights[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];p.style=doc.styles['PaperTable'];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(value)>18 else WD_ALIGN_PARAGRAPH.CENTER
   if len(data)<=7:p.paragraph_format.keep_with_next=ri<len(data)-1
   # Soft break opportunities in machine field names preserve literal content.
   p.add_run(value.replace('_','_\u200b'))
   if ri==0:
    for r in p.runs:r.bold=True
   tc=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFEFEF' if ri==0 else 'FFFFFF');tc.append(shade)
   margins=OxmlElement('w:tcMar')
   for side,v in [('top','75'),('bottom','75'),('left','85'),('right','85')]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');margins.append(el)
   tc.append(margins)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tc.append(borders)
 return t
xsl=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
eqs=[];figs=[];tables=[];current_section='摘要'
eqsplit={
 'E03':[r'C_{jk,d}(\tau)=\frac{1}{n_d}\sum_{i\in I_d}{J_{ijk}}',r'J_{ijk}=\mathbf1\{(v_{ij}\ge\tau,\ v_{ik}\le1-\tau)',r'\mathrm{or}\ (v_{ik}\ge\tau,\ v_{ij}\le1-\tau)\}'],
 'E04':[r'H^{\mathsf T}H=I_{16},\quad H^{\mathsf T}\mathbf1=0,\quad z=H^{\mathsf T}(p-\overline p)',r'\widehat L_m(p)=a_m+\beta_m^{\mathsf T}z'],
 'E08':[r'\frac{\partial L_B}{\partial N}=-\frac{\alpha U}{N},\quad\frac{\partial L_B}{\partial D}=-\frac{\beta V}{D}',r'\varepsilon_N=-\frac{\alpha U}{L_B},\quad\varepsilon_D=-\frac{\beta V}{L_B}'],
 'E09':[r'\frac{d\log D}{d\log N}=-\frac{\alpha U}{\beta V}',r'D_{\rm new}=\left(\frac{B}{L_{\rm old}-L_\infty-A N_{\rm new}^{-\alpha}}\right)^{1/\beta}'],
 'E16':[r'N_0=\left(\frac{\alpha A}{\beta B}\right)^{1/(\alpha+\beta)}K^{\beta/(\alpha+\beta)}',r'a_N=\max(N_{\min},K/D_{\max}),\quad b_N=\min(N_{\max},K/D_{\min})',r'N^*=\operatorname{clip}(N_0,a_N,b_N),\qquad D^*=K/N^*'],
 'E23':[r'z_t=\log\frac{F_t/100}{1-F_t/100}=a_F+b_Ft+e_t',r'\widehat F_{t+h}=\frac{100}{1+\exp[-(\widehat a_F+\widehat b_F(t+h))]}']}
lines=raw.splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('@TITLE '):paragraph(line[7:],'Title');continue
 if line=='@ABSTRACT':paragraph('摘要','PaperH1');continue
 if line=='@BODY':doc.add_page_break();continue
 if line=='@APPENDIX':doc.add_page_break();continue
 if line.startswith('# '):
  current_section=line[2:];paragraph(current_section,'PaperH1');continue
 if line.startswith('## '):current_section=line[3:];paragraph(current_section,'PaperH2');continue
 if line.startswith('$$ '):
  label,tex=line[3:-3].split(' | ',1);num=int(label[1:]);parts=eqsplit.get(label,[tex])
  for pi,part in enumerate(parts):
   p=paragraph('','PaperEquation');p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=pi<len(parts)-1
   mm=etree.fromstring(convert(part).encode());om=xsl(mm).getroot();p._p.append(om)
   if pi==len(parts)-1:p.add_run('    （'+str(num)+'）')
  q=1 if num<=5 else 2 if num<=12 else 3 if num<=19 else 4
  eqs.append(dict(equation_id=label,number=num,question=q,formula=tex,rendered_parts=parts,symbol_definitions='FINAL_SYMBOL_TABLE_v1.md及公式前后定义',units='N十亿参数，D十亿token；Q4分数0—100；其他见符号表',source_model={1:'Q1_MODEL_SPEC_v1 / fixed diagnostic definitions',2:'Q2_ROUND5_SPEC_v1',3:'Q3_MODEL_SPEC_v1 / Q3_COST_MODEL_v1',4:'Q4_FRONTIER_MODEL_SPEC_v1'}[q],paper_section=current_section,assumptions={1:'固定参考及局部欧氏配比关联；非质量因果效应',2:'来源内加性幂律；质量与运输仅条件情景',3:'题设成本、外生上下文、统计支持域；预算可闲置',4:'同口径评测群体；描述性回归；趋势与条件扩散'}[q]))
  continue
 if line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line);caption,path=m.groups();im=Image.open(ROOT/path)
  # Use the frozen assets verbatim; fit tall plots without tiny text.
  w=min(width_cm,12.0 if im.height/im.width>.78 else width_cm)
  p=paragraph('','PaperCaption');p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(ROOT/path),width=Cm(w))
  paragraph(caption,'PaperCaption')
  figs.append(dict(number=len(figs)+1,caption=caption,source=path,sha256=hashlib.sha256((ROOT/path).read_bytes()).hexdigest(),width_cm=w,paper_section=current_section))
  continue
 if re.match(r'^表(?:\d+|[AB]\d+) ',line):
  p=paragraph(line,'PaperCaption');p.paragraph_format.keep_with_next=True;tables.append(line);continue
 if line.startswith('|'):
  data=[x.strip() for x in line.strip('|').split('|')];data=[data]
  while i<len(lines) and lines[i].strip().startswith('|'):
   row=[x.strip() for x in lines[i].strip().strip('|').split('|')];i+=1
   if all(re.fullmatch('[-: ]+',x) for x in row):continue
   data.append(row)
  table(data);continue
 paragraph(line,'PaperRef' if re.match(r'^\[\d+\]',line) else 'PaperBody')
# Scrub editing/identity metadata; scientific AI disclosure is intentional.
for xp in ['.//w:ins','.//w:del','.//w:commentRangeStart','.//w:commentRangeEnd','.//w:commentReference']:
 for e in doc.element.xpath(xp):raise ValueError('Unexpected revision/comment in source template')
settings=doc.settings.element
for e in settings.findall(qn('w:docVars')):settings.remove(e)
for e in settings.findall(qn('w:trackRevisions')):settings.remove(e)
uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');settings.append(uf)
target=OUT/'F_final_manuscript.docx';doc.save(target)
with ZipFile(TEMPLATE) as a,ZipFile(target) as b:
 media=[p for p in a.namelist() if p.startswith('word/media/')]
 preservation={p:a.read(p)==b.read(p) for p in media}
assert all(preservation.values())
jwrite(SRC/'FINAL_EQUATION_REGISTRY_v1.json',eqs)
(SRC/'FINAL_EQUATION_REGISTRY_v1.md').write_text('# 最终公式登记\n\n'+'\n\n'.join('## 式（'+str(e['number'])+'）\n'+'\n'.join(f'- {k}: {v}' for k,v in e.items()) for e in eqs),encoding='utf-8')
jwrite(SRC/'FINAL_FIGURE_SELECTION_v1.json',figs)
jwrite(SRC/'BUILD_MANIFEST_v1.json',dict(equations=len(eqs),figures=len(figs),tables=len(tables),table_captions=tables,template_media_preservation=preservation,template_sha256=hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),source_files=source_files,render_pending=True))
(ROOT/'tmp/round8/artifact.md').write_text('''# Official template execution contract
Reference: 08_paper/round8/template_reference.docx, converted from retained official .doc.
Original .doc remains unchanged. Template render has 4 pages, one section. Its blank abstract heading leaks onto cover in LibreOffice; replace blank abstract/body slots with a section break to keep the official cover intact.
Cover preserve-only: body elements 0..6, including four logo image parts, title typography and blank identity table. Preserve the four image byte streams, positions and table structure.
Page system: A4 portrait, margins and footer distances inherited exactly from reference (top 3.0021cm, bottom 1.8486cm, left 2.2507cm, right 2.2472cm). No header; no cover page number. A new section starts abstract page number 1; body starts on next page after abstract.
Editable slots: all blank abstract/body content after body element 6. Repeating identity/event headings on abstract slots are removed in favor of official-required paper title, abstract and keywords. Body sections use 12pt Song, single spacing; title 16pt Hei; H1 14pt Hei centered. Equations are OMML, numbered 1..24. Tables repeat header rows; frozen figures are embedded unchanged.
Metadata, legacy footer and blank-page filler are editable; remove author/last editor/comments/identity. Required scientific AI disclosure is retained as body content, not hidden metadata.
QA: compare four media hashes, inspect cover and all manuscript pages, validate page numbers, OMML, captions, text bounds, references, hidden revisions and metadata. Field caching updated by rendering; no TOC is imposed by the official template.
''',encoding='utf-8')
print(json.dumps(dict(docx=str(target),equations=len(eqs),figures=len(figs),tables=len(tables)),ensure_ascii=False))
